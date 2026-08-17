from pathlib import Path
from faker import Faker
from fpdf import FPDF
fake = Faker()

# Project folders
PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"

# Create data folder if it doesn't exist
DATA_DIR.mkdir(exist_ok=True)
print("P6 test-data generator is ready!")
print("Data folder:", DATA_DIR)
print("Sample fake name:", fake.name())
# Create a simple synthetic text file
test_file = DATA_DIR / "test_01.txt"

with open(test_file, "w", encoding="utf-8") as file:
    file.write("This is a synthetic test document.\n")
    file.write(f"Name: {fake.name()}\n")
    file.write(f"City: {fake.city()}\n")

print("Created:", test_file.name)

# Create a simple synthetic PDF
pdf_file = DATA_DIR / "test_02.pdf"

pdf = FPDF()
pdf.add_page()
pdf.set_font("Arial", size=16)

pdf.cell(0, 10, "Synthetic College Notice", ln=True)

pdf.set_font("Arial", size=12)
pdf.cell(0, 10, "This is a fake document for PII detection testing.", ln=True)
pdf.cell(0, 10, f"Student Name: {fake.name()}", ln=True)
pdf.cell(0, 10, f"City: {fake.city()}", ln=True)

pdf.output(str(pdf_file))

print("Created:", pdf_file.name)

from docx import Document

# Create a clean synthetic DOCX
doc = Document()

doc.add_heading("College Notice", level=1)
doc.add_paragraph("This is a synthetic document created for PII detection testing.")
doc.add_paragraph("The college will conduct a technical workshop next week.")
doc.add_paragraph("All students are requested to attend.")

docx_file = DATA_DIR / "test_03.docx"
doc.save(docx_file)

print("Created:", docx_file.name)

# Create a synthetic PAN-like test file
pan_file = DATA_DIR / "test_04.txt"

fake_pan = "ABCDE1234F"

with open(pan_file, "w", encoding="utf-8") as file:
    file.write("Synthetic PAN detection test.\n")
    file.write(f"PAN-like value: {fake_pan}\n")
    file.write("This value is synthetic and for testing only.\n")

print("Created:", pan_file.name)

# Create a clean synthetic invoice
invoice_file = DATA_DIR / "test_05.txt"

with open(invoice_file, "w", encoding="utf-8") as file:
    file.write("INVOICE\n")
    file.write("Invoice Number: INV-2026-0045\n")
    file.write("Product: Wireless Keyboard\n")
    file.write("Quantity: 2\n")
    file.write("Total Amount: 2499 INR\n")
    file.write("Thank you for your purchase.\n")

print("Created:", invoice_file.name)


# Create a synthetic driving licence-like test file
dl_file = DATA_DIR / "test_06.txt"

with open(dl_file, "w", encoding="utf-8") as file:
    file.write("DRIVING LICENCE TEST DOCUMENT\n")
    file.write("Name: Alex Morgan\n")
    file.write("Licence Number: DL-SYN-2026-45821\n")
    file.write("Date of Birth: 15-04-2005\n")
    file.write("This is synthetic data created for PII detection testing only.\n")

print("Created:", dl_file.name)


# Create a synthetic CSV containing PII-like data
csv_file = DATA_DIR / "test_07.csv"

with open(csv_file, "w", encoding="utf-8") as file:
    file.write("name,phone,email\n")
    file.write("Alex Morgan,9876543210,alex.synthetic@example.com\n")
    file.write("Jamie Lee,9123456780,jamie.synthetic@example.com\n")

print("Created:", csv_file.name)


# Create an adversarial document with a non-PII 12-digit number
adversarial_file = DATA_DIR / "test_08.txt"

with open(adversarial_file, "w", encoding="utf-8") as file:
    file.write("INVOICE DOCUMENT\n")
    file.write("Invoice Number: 123456789012\n")
    file.write("Product: Laptop Stand\n")
    file.write("Quantity: 1\n")
    file.write("Amount: 1499 INR\n")
    file.write("This number is an invoice identifier, not government-issued PII.\n")

print("Created:", adversarial_file.name)


import json

# Create a synthetic JSON file containing PII
json_file = DATA_DIR / "test_09.json"

data = {
    "name": "Taylor Morgan",
    "phone": "9000012345",
    "email": "taylor.synthetic@example.com",
    "city": "Hyderabad"
}

with open(json_file, "w", encoding="utf-8") as file:
    json.dump(data, file, indent=4)

print("Created:", json_file.name)